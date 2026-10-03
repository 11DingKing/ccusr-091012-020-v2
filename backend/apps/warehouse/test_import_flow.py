"""品种目录暂存导入：逐行校验、修订、原子发布、并发与追溯的端到端测试。"""
import io
import threading
from datetime import timedelta
from unittest.mock import patch

from django.core.files.uploadedfile import SimpleUploadedFile
from django.db import IntegrityError, transaction
from django.test import TransactionTestCase
from django.utils import timezone
from openpyxl import Workbook, load_workbook
from rest_framework.test import APIClient

from apps.authentication.backends import generate_token
from apps.authentication.models import User
from apps.warehouse.imports.publishing import publish_batch
from apps.warehouse.models import (
    Category, ImportBatch, ImportRow, Unit, Variety,
)


def build_xlsx(rows):
    """rows: [(品种, 品类, 单位), ...] → Excel 字节"""
    wb = Workbook()
    ws = wb.active
    ws.append(['品种', '品类', '单位'])
    for row in rows:
        ws.append(list(row))
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


def upload_file(client, rows, filename='varieties.xlsx'):
    content = build_xlsx(rows)
    return client.post(
        '/api/varieties/import/',
        {'file': SimpleUploadedFile(filename, content,
                                    content_type='application/vnd.openxmlformats-officedocument.spreadsheetml.sheet')},
        format='multipart',
    )


class ImportFixture(TransactionTestCase):
    def setUp(self):
        self.user = User.objects.create_user('import-user', 'testpass123', role='admin')
        self.client = APIClient()
        self.client.credentials(HTTP_AUTHORIZATION=f'Bearer {generate_token(self.user)}')
        self.unit = Unit.objects.create(name='件', created_by=self.user)
        self.category = Category.objects.create(name='受控器材', unit=self.unit, created_by=self.user)
        self.other_unit = Unit.objects.create(name='台', created_by=self.user)
        self.other_category = Category.objects.create(name='通信设备', unit=self.other_unit, created_by=self.user)
        # 正式表中已有的品种，用于“更新”命中
        self.existing = Variety.objects.create(name='记录终端', category=self.category, created_by=self.user)

    def batch_url(self, batch_id, action=''):
        base = f'/api/varieties/import/batches/{batch_id}/'
        return base + (f'{action}/' if action else '')


class StagingAndValidationTest(ImportFixture):
    def test_upload_only_stages_and_never_writes_formal_table(self):
        before = set(Variety.objects.values_list('id', flat=True))
        resp = upload_file(self.client, [
            ('执法记录仪', '受控器材', '件'),      # 新增
            ('记录终端', '受控器材', '件'),        # 更新
            ('未知品类机', '不存在品类', '件'),    # 错误：品类不存在
            ('单位错配', '受控器材', '台'),        # 错误：单位不匹配
        ])
        self.assertEqual(resp.status_code, 200, resp.content)
        data = resp.json()['data']
        self.assertEqual(data['status'], ImportBatch.STATUS_DRAFT)
        self.assertEqual(data['total_count'], 4)
        self.assertEqual(data['create_count'], 1)
        self.assertEqual(data['update_count'], 1)
        self.assertEqual(data['error_count'], 2)
        # 正式表没有任何新增
        self.assertEqual(set(Variety.objects.values_list('id', flat=True)), before)

        results = {r['sheet_row']: r for r in data['rows']}
        self.assertEqual(results[2]['result'], ImportRow.RESULT_CREATE)
        self.assertEqual(results[3]['result'], ImportRow.RESULT_UPDATE)
        self.assertEqual(results[3]['variety'], self.existing.id)
        self.assertTrue(results[4]['is_error'])
        self.assertIn('不存在', results[4]['error_message'])
        self.assertIn('不匹配', results[5]['error_message'])

        # 影响摘要
        summary = data['summary']
        self.assertEqual(summary['create_count'], 1)
        self.assertEqual(summary['update_count'], 1)
        self.assertEqual(len(summary['creates']), 1)
        self.assertEqual(len(summary['updates']), 1)

    def test_duplicate_file_is_rejected(self):
        rows = [('执法记录仪', '受控器材', '件')]
        first = upload_file(self.client, rows)
        self.assertEqual(first.status_code, 200)
        second = upload_file(self.client, rows)
        self.assertEqual(second.status_code, 409)
        self.assertEqual(
            second.json()['data']['existing_batch_id'],
            first.json()['data']['id'],
        )
        # 不同文件名、相同内容也算重复
        third = upload_file(self.client, rows, filename='copy.xlsx')
        self.assertEqual(third.status_code, 409)

    def test_same_business_key_multiple_rows_is_conflict(self):
        resp = upload_file(self.client, [
            ('重复品种', '受控器材', '件'),
            ('重复品种', '受控器材', '件'),
        ])
        self.assertEqual(resp.status_code, 200)
        data = resp.json()['data']
        self.assertEqual(data['conflict_count'], 1)
        rows = {r['sheet_row']: r for r in data['rows']}
        self.assertEqual(rows[2]['result'], ImportRow.RESULT_CREATE)
        self.assertEqual(rows[3]['result'], ImportRow.RESULT_CONFLICT)
        self.assertIn('第2行', rows[3]['error_message'])
        # 同名品种在不同品类下不是冲突（且文件指纹不同）
        resp2 = upload_file(self.client, [
            ('跨品类同名', '受控器材', '件'),
            ('跨品类同名', '通信设备', '台'),
        ], filename='distinct.xlsx')
        self.assertEqual(resp2.status_code, 200)
        self.assertEqual(resp2.json()['data']['conflict_count'], 0)

    def test_bad_file_and_empty_file(self):
        bad = self.client.post(
            '/api/varieties/import/',
            {'file': SimpleUploadedFile('bad.xlsx', b'not an excel file')},
            format='multipart',
        )
        self.assertEqual(bad.status_code, 400)
        self.assertEqual(ImportBatch.objects.count(), 0)

        empty = self.client.post(
            '/api/varieties/import/',
            {'file': SimpleUploadedFile('empty.xlsx', build_xlsx([]),
                                        content_type='application/vnd.openxmlformats-officedocument.spreadsheetml.sheet')},
            format='multipart',
        )
        self.assertEqual(empty.status_code, 400)

    def test_requires_authentication(self):
        resp = APIClient().get('/api/varieties/import/batches/')
        self.assertEqual(resp.status_code, 401)


class RowRevisionTest(ImportFixture):
    def test_fix_only_error_row_without_reupload(self):
        resp = upload_file(self.client, [
            ('执法记录仪', '受控器材', '件'),
            ('坏数据', '不存在品类', '件'),
        ])
        batch = resp.json()['data']
        self.assertEqual(batch['status'], ImportBatch.STATUS_DRAFT)
        bad_row = next(r for r in batch['rows'] if r['sheet_row'] == 3)

        fixed = self.client.patch(
            f"/api/varieties/import/batches/{batch['id']}/rows/{bad_row['id']}/",
            {'variety_name': '修好的品种', 'category_name': '受控器材', 'unit_name': '件'},
            format='json',
        )
        self.assertEqual(fixed.status_code, 200, fixed.content)
        data = fixed.json()['data']
        self.assertEqual(data['status'], ImportBatch.STATUS_READY)
        self.assertEqual(data['error_count'], 0)
        self.assertEqual(data['create_count'], 2)
        # 暂存行内容确实被修订
        row3 = next(r for r in data['rows'] if r['sheet_row'] == 3)
        self.assertEqual(row3['variety_name'], '修好的品种')
        # 正式表仍未写入
        self.assertFalse(Variety.objects.filter(name='修好的品种').exists())

    def test_delete_duplicate_row_clears_conflict(self):
        resp = upload_file(self.client, [
            ('重复品种', '受控器材', '件'),
            ('重复品种', '受控器材', '件'),
        ])
        batch = resp.json()['data']
        dup_row = next(r for r in batch['rows'] if r['result'] == ImportRow.RESULT_CONFLICT)

        deleted = self.client.delete(
            f"/api/varieties/import/batches/{batch['id']}/rows/{dup_row['id']}/delete/"
        )
        self.assertEqual(deleted.status_code, 200)
        data = deleted.json()['data']
        self.assertEqual(data['total_count'], 1)
        self.assertEqual(data['conflict_count'], 0)
        self.assertEqual(data['status'], ImportBatch.STATUS_READY)

    def test_cannot_fix_published_batch(self):
        resp = upload_file(self.client, [('执法记录仪', '受控器材', '件')])
        batch_id = resp.json()['data']['id']
        self.client.post(f'/api/varieties/import/batches/{batch_id}/publish/')
        row_id = ImportRow.objects.get(batch_id=batch_id).id

        fix = self.client.patch(
            f'/api/varieties/import/batches/{batch_id}/rows/{row_id}/',
            {'variety_name': 'x', 'category_name': '受控器材', 'unit_name': '件'},
            format='json',
        )
        self.assertEqual(fix.status_code, 400)

    def test_revalidate_picks_up_formal_table_change(self):
        resp = upload_file(self.client, [('后到的品种', '受控器材', '件')])
        batch = resp.json()['data']
        row = batch['rows'][0]
        self.assertEqual(row['result'], ImportRow.RESULT_CREATE)

        # 上传后、发布前，正式表被别人写入了同一业务键
        Variety.objects.create(name='后到的品种', category=self.category, created_by=self.user)

        revalidated = self.client.post(
            f"/api/varieties/import/batches/{batch['id']}/revalidate/"
        )
        self.assertEqual(revalidated.status_code, 200)
        self.assertEqual(revalidated.json()['data']['rows'][0]['result'], ImportRow.RESULT_UPDATE)


class PublishTest(ImportFixture):
    def test_publish_atomically_creates_and_updates_with_traceability(self):
        resp = upload_file(self.client, [
            ('执法记录仪', '受控器材', '件'),
            ('对讲机', '通信设备', '台'),
            ('记录终端', '受控器材', '件'),
        ])
        batch_id = resp.json()['data']['id']

        publish = self.client.post(f'/api/varieties/import/batches/{batch_id}/publish/')
        self.assertEqual(publish.status_code, 200, publish.content)
        data = publish.json()['data']
        self.assertEqual(data['status'], ImportBatch.STATUS_PUBLISHED)
        self.assertEqual(data['create_count'], 2)
        self.assertEqual(data['update_count'], 1)
        self.assertEqual(data['published_by'], self.user.id)
        self.assertIsNotNone(data['published_at'])

        # 正式表结果
        self.assertTrue(Variety.objects.filter(name='执法记录仪', category=self.category).exists())
        self.assertTrue(Variety.objects.filter(name='对讲机', category=self.other_category).exists())
        self.assertEqual(Variety.objects.count(), 3)  # 初始1 + 新增2，更新不新增

        # 逐行可追溯：每一行都有终态结果和正式品种引用
        rows = ImportRow.objects.filter(batch_id=batch_id).order_by('sheet_row')
        self.assertEqual([r.result for r in rows],
                         [ImportRow.RESULT_CREATED, ImportRow.RESULT_CREATED, ImportRow.RESULT_UPDATED])
        for row in rows:
            self.assertIsNotNone(row.published_variety_id)
        self.assertEqual(rows[2].published_variety_id, self.existing.id)

    def test_publish_blocked_while_errors_exist_and_writes_nothing(self):
        resp = upload_file(self.client, [
            ('执法记录仪', '受控器材', '件'),
            ('坏数据', '不存在品类', '件'),
        ])
        batch_id = resp.json()['data']['id']
        before_count = Variety.objects.count()

        publish = self.client.post(f'/api/varieties/import/batches/{batch_id}/publish/')
        self.assertEqual(publish.status_code, 400)
        # 即便第一行合法，也没有任何正式数据落库
        self.assertEqual(Variety.objects.count(), before_count)
        self.assertFalse(Variety.objects.filter(name='执法记录仪').exists())
        # 批次仍开放可修订
        self.assertEqual(ImportBatch.objects.get(pk=batch_id).status, ImportBatch.STATUS_DRAFT)

    def test_duplicate_publish_is_rejected(self):
        resp = upload_file(self.client, [('执法记录仪', '受控器材', '件')])
        batch_id = resp.json()['data']['id']
        first = self.client.post(f'/api/varieties/import/batches/{batch_id}/publish/')
        self.assertEqual(first.status_code, 200)
        second = self.client.post(f'/api/varieties/import/batches/{batch_id}/publish/')
        self.assertEqual(second.status_code, 400)
        self.assertEqual(Variety.objects.filter(name='执法记录仪').count(), 1)

    def test_mid_publish_failure_rolls_back_everything(self):
        resp = upload_file(self.client, [
            ('执法记录仪', '受控器材', '件'),
            ('对讲机', '通信设备', '台'),
        ])
        batch_id = resp.json()['data']['id']

        # 模拟发布写入途中异常
        with patch.object(Variety.objects, 'bulk_create', side_effect=RuntimeError('磁盘故障')):
            publish = self.client.post(f'/api/varieties/import/batches/{batch_id}/publish/')
        self.assertEqual(publish.status_code, 400)
        self.assertIn('回滚', publish.json()['message'])

        # 正式表零残留
        self.assertFalse(Variety.objects.filter(name__in=['执法记录仪', '对讲机']).exists())
        batch = ImportBatch.objects.get(pk=batch_id)
        self.assertIn(batch.status, [ImportBatch.STATUS_DRAFT, ImportBatch.STATUS_READY])
        self.assertTrue(batch.is_open)
        self.assertIn('磁盘故障', batch.publish_error)
        # 行结果被重新校验恢复，仍可再次发布
        self.assertEqual(
            ImportRow.objects.filter(batch_id=batch_id, result=ImportRow.RESULT_CREATE).count(), 2
        )

        # 故障恢复后可以正常发布
        retry = self.client.post(f'/api/varieties/import/batches/{batch_id}/publish/')
        self.assertEqual(retry.status_code, 200)
        self.assertEqual(Variety.objects.filter(name__in=['执法记录仪', '对讲机']).count(), 2)

    def test_integrity_error_during_publish_rolls_back(self):
        resp = upload_file(self.client, [('执法记录仪', '受控器材', '件')])
        batch_id = resp.json()['data']['id']

        with patch.object(
            Variety.objects, 'bulk_create',
            side_effect=IntegrityError('unique constraint violated'),
        ):
            publish = self.client.post(f'/api/varieties/import/batches/{batch_id}/publish/')
        self.assertEqual(publish.status_code, 400)
        self.assertFalse(Variety.objects.filter(name='执法记录仪').exists())
        self.assertTrue(ImportBatch.objects.get(pk=batch_id).is_open)

    def test_discard_then_same_file_can_be_uploaded_again(self):
        rows = [('执法记录仪', '受控器材', '件')]
        resp = upload_file(self.client, rows)
        batch_id = resp.json()['data']['id']

        discarded = self.client.post(f'/api/varieties/import/batches/{batch_id}/discard/')
        self.assertEqual(discarded.status_code, 200)
        self.assertEqual(ImportRow.objects.filter(batch_id=batch_id).count(), 0)

        re_upload = upload_file(self.client, rows)
        self.assertEqual(re_upload.status_code, 200)

    def test_result_export_traces_every_row(self):
        resp = upload_file(self.client, [
            ('执法记录仪', '受控器材', '件'),
            ('记录终端', '受控器材', '件'),
        ])
        batch_id = resp.json()['data']['id']
        self.client.post(f'/api/varieties/import/batches/{batch_id}/publish/')

        export = self.client.get(f'/api/varieties/import/batches/{batch_id}/result/')
        self.assertEqual(export.status_code, 200)
        wb = load_workbook(io.BytesIO(export.content))
        ws = wb.active
        header = [c.value for c in ws[1]]
        self.assertEqual(header, ['Excel行号', '品种', '品类', '单位', '结果', '错误原因', '正式品种ID'])
        body = [[c.value for c in row] for row in ws.iter_rows(min_row=2)]
        self.assertEqual(body[0][4], '已新增')
        self.assertEqual(body[1][4], '已更新')
        self.assertTrue(body[0][6])  # 正式品种ID可追溯

    def test_batch_listing_and_detail(self):
        upload_file(self.client, [('执法记录仪', '受控器材', '件')])
        listing = self.client.get('/api/varieties/import/batches/')
        self.assertEqual(listing.status_code, 200)
        self.assertEqual(listing.json()['data']['total'], 1)

        batch_id = listing.json()['data']['list'][0]['id']
        detail = self.client.get(f'/api/varieties/import/batches/{batch_id}/')
        self.assertEqual(detail.status_code, 200)
        self.assertEqual(len(detail.json()['data']['rows']), 1)


class StalePublishingRecoveryTest(ImportFixture):
    def test_stale_publishing_status_is_reclaimable(self):
        resp = upload_file(self.client, [('执法记录仪', '受控器材', '件')])
        batch_id = resp.json()['data']['id']

        # 模拟进程崩溃：状态停留在 publishing，开始时间在很久以前
        with transaction.atomic():
            ImportBatch.objects.filter(pk=batch_id).update(
                status=ImportBatch.STATUS_PUBLISHING,
                publish_started_at=timezone.now() - timedelta(hours=1),
            )

        retry = self.client.post(f'/api/varieties/import/batches/{batch_id}/publish/')
        self.assertEqual(retry.status_code, 200)
        self.assertEqual(ImportBatch.objects.get(pk=batch_id).status, ImportBatch.STATUS_PUBLISHED)


class ConcurrentPublishTest(ImportFixture):
    """同一批次并发发布：恰好一个成功，另一个被拒绝，正式表无重复无残留。

    直接调用服务层（不经过视图的预读），以聚焦“发布认领”的串行化语义；
    SQLite 共享缓存下等待锁会立即抛 locked，生产文件库则由 busy timeout 等待，
    两种情况下第二个发布都被明确拒绝，最终状态一致。
    """

    def test_two_threads_publish_same_batch(self):
        resp = upload_file(self.client, [
            (f'并发品种{i}', '受控器材', '件') for i in range(5)
        ])
        batch_id = resp.json()['data']['id']
        barrier = threading.Barrier(2)
        outcomes = []
        lock = threading.Lock()

        def worker():
            barrier.wait()
            try:
                publish_batch(batch_id, self.user)
                code = 200
            except Exception as exc:  # noqa: BLE001
                code = getattr(exc, 'status_code', 500)
            with lock:
                outcomes.append(code)

        t1 = threading.Thread(target=worker)
        t2 = threading.Thread(target=worker)
        t1.start(); t2.start()
        t1.join(15); t2.join(15)

        self.assertEqual(outcomes.count(200), 1, outcomes)
        self.assertTrue(
            all(c in (400, 409) for c in outcomes if c != 200), outcomes
        )
        self.assertEqual(
            Variety.objects.filter(name__startswith='并发品种').count(), 5
        )
        self.assertEqual(
            ImportBatch.objects.get(pk=batch_id).status,
            ImportBatch.STATUS_PUBLISHED,
        )
        # 每一行恰好写入一个正式品种，可追溯
        rows = ImportRow.objects.filter(batch_id=batch_id)
        self.assertEqual(rows.filter(result=ImportRow.RESULT_CREATED).count(), 5)
        self.assertEqual(rows.exclude(published_variety__isnull=True).count(), 5)


class ServiceLevelConcurrencyTest(ImportFixture):
    def test_in_flight_publish_blocks_a_second_claim(self):
        resp = upload_file(self.client, [('执法记录仪', '受控器材', '件')])
        batch_id = resp.json()['data']['id']

        # 模拟已有发布事务正在进行：状态停留在 publishing 且未超时
        ImportBatch.objects.filter(pk=batch_id).update(
            status=ImportBatch.STATUS_PUBLISHING,
            publish_started_at=timezone.now(),
        )

        from apps.warehouse.imports.publishing import PublishError
        with self.assertRaises(PublishError) as ctx:
            publish_batch(batch_id, self.user)
        self.assertIn('发布', ctx.exception.message)
        self.assertEqual(ctx.exception.status_code, 409)
        # 被拒方没有写入任何正式数据，批次仍由先到的发布持有
        self.assertFalse(Variety.objects.filter(name='执法记录仪').exists())
        self.assertEqual(
            ImportBatch.objects.get(pk=batch_id).status,
            ImportBatch.STATUS_PUBLISHING,
        )

    def test_published_batch_claim_is_rejected(self):
        resp = upload_file(self.client, [('执法记录仪', '受控器材', '件')])
        batch_id = resp.json()['data']['id']
        publish_batch(batch_id, self.user)

        from apps.warehouse.imports.publishing import PublishError
        with self.assertRaises(PublishError) as ctx:
            publish_batch(batch_id, self.user)
        self.assertIn('已发布', ctx.exception.message)
        self.assertEqual(Variety.objects.filter(name='执法记录仪').count(), 1)
