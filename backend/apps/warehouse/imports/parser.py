"""Excel 文件解析：把上传文件解析为规范化的原始行，并计算文件指纹。"""
import hashlib
import io

from openpyxl import load_workbook

# 模板列顺序：品种 / 品类 / 单位（与 VarietyTemplateView 保持一致）
COLUMNS = ('variety', 'category', 'unit')


class ImportFileError(Exception):
    """文件无法解析或内容为空"""


def read_uploaded_file(uploaded_file):
    """读取上传文件的二进制内容并计算 SHA256 指纹。"""
    uploaded_file.seek(0)
    content = uploaded_file.read()
    uploaded_file.seek(0)
    file_hash = hashlib.sha256(content).hexdigest()
    return content, file_hash


def _cell_text(value):
    if value is None:
        return ''
    return str(value).strip()


def parse_rows(content):
    """解析 Excel 字节内容，返回 [{sheet_row, variety, category, unit}, ...]。

    - 第一行为表头；
    - 三个字段全部为空的行视为空白行，跳过；
    - 任一字段有值即视为数据行，缺失字段交由校验阶段报错。
    """
    try:
        wb = load_workbook(io.BytesIO(content), read_only=True, data_only=True)
    except Exception as exc:  # openpyxl 对坏文件会抛各种异常
        raise ImportFileError('文件格式错误，请上传Excel文件') from exc

    ws = wb.active
    rows = []
    for sheet_row, row in enumerate(ws.iter_rows(min_row=2, values_only=True), start=2):
        values = list(row) + [None] * max(0, 3 - len(row))
        variety, category, unit = (_cell_text(v) for v in values[:3])
        if not any((variety, category, unit)):
            continue
        rows.append({
            'sheet_row': sheet_row,
            'variety': variety,
            'category': category,
            'unit': unit,
        })
    wb.close()

    if not rows:
        raise ImportFileError('文件中没有可导入的数据行')
    return rows
