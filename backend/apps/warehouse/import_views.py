"""品种目录导入接口：暂存上传、逐行校验、修订与原子发布。"""
import io
import logging

from django.http import HttpResponse
from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from rest_framework.parsers import FormParser, MultiPartParser
from rest_framework.permissions import IsAuthenticated
from rest_framework.views import APIView

from apps.core.response import error_response, success_response

from .import_serializers import (
    ImportBatchDetailSerializer,
    ImportBatchListSerializer,
    ImportRowFixSerializer,
)
from .imports.parser import ImportFileError
from .imports.publishing import PublishError, publish_batch
from .imports.staging import (
    StagingError,
    create_batch,
    delete_row,
    discard_batch,
    fix_row,
)
from .imports.validation import validate_batch
from .models import ImportBatch

logger = logging.getLogger('apps')


def _get_batch(pk):
    return ImportBatch.objects.filter(pk=pk).first()


def _get_row(batch, row_pk):
    return batch.rows.filter(pk=row_pk).first()


class VarietyImportUploadView(APIView):
    """上传品种目录文件：只暂存并逐行校验，不写正式表"""
    permission_classes = [IsAuthenticated]
    parser_classes = [MultiPartParser, FormParser]

    def post(self, request):
        uploaded = request.FILES.get('file')
        if uploaded is None:
            return error_response(message='请上传文件')

        try:
            batch = create_batch(uploaded, request.user)
        except ImportFileError as exc:
            return error_response(message=str(exc))
        except StagingError as exc:
            data = None
            if exc.existing_batch is not None:
                data = {'existing_batch_id': exc.existing_batch.id}
            return error_response(message=exc.message, code=exc.status_code, data=data)

        serializer = ImportBatchDetailSerializer(batch)
        return success_response(
            data=serializer.data,
            message='文件已暂存并完成逐行校验，修订错误行后即可发布',
        )


class VarietyImportListView(APIView):
    """导入版本列表"""
    permission_classes = [IsAuthenticated]

    def get(self, request):
        queryset = ImportBatch.objects.all()
        status = request.query_params.get('status')
        if status:
            queryset = queryset.filter(status=status)

        total = queryset.count()
        try:
            page = max(1, int(request.query_params.get('page', 1)))
            page_size = max(1, int(request.query_params.get('page_size', 10)))
        except (TypeError, ValueError):
            page, page_size = 1, 10

        batches = queryset[(page - 1) * page_size:page * page_size]
        return success_response(data={
            'list': ImportBatchListSerializer(batches, many=True).data,
            'total': total,
            'page': page,
            'page_size': page_size,
        })


class VarietyImportDetailView(APIView):
    """导入版本详情：影响摘要 + 全部暂存行（发布后含每行发布结果）"""
    permission_classes = [IsAuthenticated]

    def get(self, request, pk):
        batch = _get_batch(pk)
        if batch is None:
            return error_response(message='导入批次不存在', code=404)
        return success_response(data=ImportBatchDetailSerializer(batch).data)


class VarietyImportRowFixView(APIView):
    """修订单条错误/冲突行，无需重传整份文件"""
    permission_classes = [IsAuthenticated]

    def patch(self, request, pk, row_pk):
        batch = _get_batch(pk)
        if batch is None:
            return error_response(message='导入批次不存在', code=404)
        row = _get_row(batch, row_pk)
        if row is None:
            return error_response(message='该行不属于该导入批次或不存在', code=404)

        serializer = ImportRowFixSerializer(data=request.data)
        if not serializer.is_valid():
            first_error = next(iter(serializer.errors.values()))[0]
            return error_response(message=str(first_error))

        try:
            fix_row(
                batch, row,
                serializer.validated_data['variety_name'].strip(),
                serializer.validated_data['category_name'].strip(),
                serializer.validated_data['unit_name'].strip(),
            )
        except StagingError as exc:
            return error_response(message=exc.message, code=exc.status_code)

        batch = _get_batch(pk)
        return success_response(
            data=ImportBatchDetailSerializer(batch).data,
            message='该行已修订，批次已重新校验',
        )


class VarietyImportRowDeleteView(APIView):
    """删除单条暂存行（用于消除文件内重复业务键）"""
    permission_classes = [IsAuthenticated]

    def delete(self, request, pk, row_pk):
        batch = _get_batch(pk)
        if batch is None:
            return error_response(message='导入批次不存在', code=404)
        row = _get_row(batch, row_pk)
        if row is None:
            return error_response(message='该行不属于该导入批次或不存在', code=404)

        try:
            delete_row(batch, row)
        except StagingError as exc:
            return error_response(message=exc.message, code=exc.status_code)

        batch = _get_batch(pk)
        return success_response(
            data=ImportBatchDetailSerializer(batch).data,
            message='该行已删除，批次已重新校验',
        )


class VarietyImportRevalidateView(APIView):
    """对暂存批次按当前正式表重新校验"""
    permission_classes = [IsAuthenticated]

    def post(self, request, pk):
        batch = _get_batch(pk)
        if batch is None:
            return error_response(message='导入批次不存在', code=404)
        if not batch.is_open:
            return error_response(
                message='该批次已发布，不能重新校验',
                code=400,
            )
        validate_batch(batch)
        return success_response(
            data=ImportBatchDetailSerializer(batch).data,
            message='重新校验完成',
        )


class VarietyImportPublishView(APIView):
    """原子发布：全部行一次性写入正式表，失败整体回滚"""
    permission_classes = [IsAuthenticated]

    def post(self, request, pk):
        batch = _get_batch(pk)
        if batch is None:
            return error_response(message='导入批次不存在', code=404)

        try:
            batch = publish_batch(pk, request.user)
        except PublishError as exc:
            logger.info(
                'User %s publish import batch %s rejected: %s',
                request.user.id, pk, exc.message,
            )
            return error_response(message=exc.message, code=exc.status_code)

        logger.info(
            'User %s published import batch %s: +%s ~%s',
            request.user.id, pk, batch.create_count, batch.update_count,
        )
        return success_response(
            data=ImportBatchDetailSerializer(batch).data,
            message=(
                f'发布成功：新增 {batch.create_count} 个，'
                f'更新 {batch.update_count} 个'
            ),
        )


class VarietyImportDiscardView(APIView):
    """丢弃未发布的导入版本（物理删除暂存数据，正式表不受影响）"""
    permission_classes = [IsAuthenticated]

    def post(self, request, pk):
        batch = _get_batch(pk)
        if batch is None:
            return error_response(message='导入批次不存在', code=404)
        try:
            discard_batch(batch)
        except StagingError as exc:
            return error_response(message=exc.message, code=exc.status_code)
        return success_response(message='导入批次已丢弃，正式数据未受影响')


class VarietyImportResultExportView(APIView):
    """导出导入版本的逐行结果，供合作单位离线核对与归档"""
    permission_classes = [IsAuthenticated]

    def get(self, request, pk):
        batch = _get_batch(pk)
        if batch is None:
            return error_response(message='导入批次不存在', code=404)

        wb = Workbook()
        ws = wb.active
        ws.title = '导入结果'

        header_font = Font(bold=True, color='FFFFFF')
        header_fill = PatternFill(start_color='4F46E5', end_color='4F46E5', fill_type='solid')
        center = Alignment(horizontal='center', vertical='center')
        border = Border(
            left=Side(style='thin'), right=Side(style='thin'),
            top=Side(style='thin'), bottom=Side(style='thin'),
        )

        headers = ['Excel行号', '品种', '品类', '单位', '结果', '错误原因', '正式品种ID']
        for col, title in enumerate(headers, 1):
            cell = ws.cell(row=1, column=col, value=title)
            cell.font = header_font
            cell.fill = header_fill
            cell.alignment = center
            cell.border = border

        for idx, row in enumerate(batch.rows.order_by('sheet_row'), start=2):
            values = [
                row.sheet_row,
                row.variety_name,
                row.category_name,
                row.unit_name,
                row.get_result_display(),
                row.error_message,
                row.published_variety_id or '',
            ]
            for col, value in enumerate(values, 1):
                ws.cell(row=idx, column=col, value=value).border = border

        widths = [10, 25, 20, 12, 10, 40, 12]
        for col, width in enumerate(widths, start=1):
            ws.column_dimensions[chr(64 + col)].width = width

        output = io.BytesIO()
        wb.save(output)
        output.seek(0)
        response = HttpResponse(
            output.read(),
            content_type='application/vnd.openxmlformats-officedocument.spreadsheetml.sheet',
        )
        response['Content-Disposition'] = (
            f'attachment; filename=import_batch_{batch.id}_result.xlsx'
        )
        return response
